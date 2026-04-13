# Google Drive service — reads compliance documents from a shared Drive folder
#
# Additional dependencies required:
#   pip install PyPDF2 openpyxl python-docx
#
# These are also listed in requirements.txt

import io
import os

import PyPDF2
import openpyxl
from docx import Document
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

from config.settings import GOOGLE_DRIVE_FOLDER_ID

# Path to the service account credentials file in the project root
CREDENTIALS_FILE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "credentials.json"
)

# Read-only Drive scope — the service account only needs to list and download files
SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]


def get_drive_service():
    """
    Authenticates using the service account credentials.json and returns an
    authorised Google Drive API v3 service object.
    """
    creds = Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=SCOPES)
    service = build("drive", "v3", credentials=creds)
    return service


def list_files():
    """
    Lists all files inside the configured GOOGLE_DRIVE_FOLDER_ID folder.
    Returns a list of dicts, each with keys: id, name, mimeType.
    Returns an empty list if the folder is empty or an error occurs.
    """
    try:
        service = get_drive_service()

        # Query Drive for files whose parent is the target folder and are not trashed
        results = service.files().list(
            q=f"'{GOOGLE_DRIVE_FOLDER_ID}' in parents and trashed = false",
            fields="files(id, name, mimeType)"
        ).execute()

        files = results.get("files", [])
        return files

    except Exception as e:
        print(f"[drive_service] Error listing files: {e}")
        return []


def read_file(file_id, mime_type):
    """
    Downloads a file from Drive by its file_id and extracts its text content
    based on the mime_type.

    Supported formats:
        - PDF  (application/pdf)                                          → PyPDF2
        - Excel (.xlsx)                                                   → openpyxl
        - Word  (.docx)                                                   → python-docx
        - Plain text (text/plain)                                         → raw decode

    Returns the extracted text as a string, or an empty string on failure.
    """
    try:
        service = get_drive_service()

        # Download the raw file bytes into an in-memory buffer
        request = service.files().get_media(fileId=file_id)
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)

        done = False
        while not done:
            _, done = downloader.next_chunk()

        buffer.seek(0)  # Rewind the buffer before reading

        # ── PDF ─────────────────────────────────────────────────────────────
        if mime_type == "application/pdf":
            reader = PyPDF2.PdfReader(buffer)
            text = ""
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
            return text.strip()

        # ── Excel (.xlsx) ────────────────────────────────────────────────────
        elif mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            workbook = openpyxl.load_workbook(buffer, data_only=True)
            lines = []
            for sheet in workbook.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    # Filter out completely empty rows
                    row_values = [str(cell) for cell in row if cell is not None]
                    if row_values:
                        lines.append("\t".join(row_values))
            return "\n".join(lines).strip()

        # ── Word (.docx) ─────────────────────────────────────────────────────
        elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = Document(buffer)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs).strip()

        # ── Plain text ───────────────────────────────────────────────────────
        elif mime_type == "text/plain":
            return buffer.read().decode("utf-8", errors="ignore").strip()

        else:
            print(f"[drive_service] Unsupported mime type: {mime_type}")
            return ""

    except Exception as e:
        print(f"[drive_service] Error reading file {file_id}: {e}")
        return ""


# MIME types treated as images — passed to AI vision rather than parsed as text
IMAGE_MIME_TYPES = {
    "image/jpeg",
    "image/jpg",
    "image/png",
    "image/webp",
}

# Module-level cache — populated on first fetch, reused on every subsequent call.
# Call refresh_cache() to force a re-fetch after documents are updated in Drive.
_cache = {"text": None, "image_ids": None}


def get_all_documents_text():
    """
    Returns all compliance document text and image file IDs from the Drive folder.

    On the first call the documents are fetched from Drive and stored in the
    module-level _cache. Every subsequent call returns the cached copy instantly
    without hitting the Drive API again.

    Returns:
        (combined_text: str, image_file_ids: list[str])
    """
    # Return cached result if available
    if _cache["text"] is not None:
        print("[drive_service] Using cached documents.")
        return _cache["text"], _cache["image_ids"]

    print("[drive_service] Fetching fresh documents from Drive...")

    files = list_files()

    if not files:
        return "", []

    sections = []
    image_file_ids = []

    for file in files:
        file_id = file["id"]
        file_name = file["name"]
        mime_type = file["mimeType"]

        print(f"[drive_service] Reading: {file_name} ({mime_type})")

        if mime_type in IMAGE_MIME_TYPES:
            # Queue the image for AI vision and insert a placeholder in the text
            image_file_ids.append(file_id)
            sections.append(
                f"=== [IMAGE FILE: {file_name}] === "
                "(image content will be analyzed by AI vision)"
            )
            continue

        text = read_file(file_id, mime_type)

        if text:
            # Use the file name as a section header so the AI knows the source
            sections.append(f"=== {file_name} ===\n{text}")
        else:
            print(f"[drive_service] Skipping {file_name} — no text extracted.")

    combined_text = "\n\n".join(sections)

    # Store in cache for all future calls
    _cache["text"] = combined_text
    _cache["image_ids"] = image_file_ids

    return combined_text, image_file_ids


def refresh_cache():
    """
    Clears the document cache and immediately re-fetches all files from Drive.
    Call this after Dekel updates any compliance documents in the Drive folder
    so the bot picks up the changes without needing a restart.

    Returns:
        (combined_text: str, image_file_ids: list[str])
    """
    print("[drive_service] Refreshing document cache...")
    _cache["text"] = None
    _cache["image_ids"] = None
    return get_all_documents_text()
